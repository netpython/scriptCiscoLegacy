#!/usr/bin/env python3
"""Audit multi-équipements IOS : sécurité, services et conformité globale."""
from __future__ import annotations
import re, sys
from collections import Counter
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'scripts'))
from common import devices, parallel, parser, run, stamp, OUTPUT, setup
RULES = {'SSH_V2': ('(?m)^ip ssh version 2$', 10), 'NO_HTTP': ('(?m)^no ip http server$', 10), 'AAA': ('(?m)^aaa new-model$', 15), 'NTP': ('(?m)^ntp server ', 10), 'SYSLOG': ('(?m)^logging host ', 10), 'SNMP_SECURE': ('(?m)^snmp-server group .* v3', 15), 'BANNER': ('(?m)^banner (motd|login)', 5), 'PASSWORD_ENCRYPTION': ('(?m)^service password-encryption$', 5), 'NO_CDP_GLOBAL': ('(?m)^no cdp run$', 5), 'VTY_SSH': ('(?ms)^line vty .*?transport input ssh', 15)}
log = setup('enterprise_compliance')
p = parser('Audit conformité Cisco IOS avancé')
args = p.parse_args()

def audit(d):
    cfg = run(d, 'show running-config')
    row = {'device': d.get('name', d['host']), 'host': d['host']}
    earned = 0
    for name, (rx, weight) in RULES.items():
        ok = bool(re.search(rx, cfg))
        row[name] = 'PASS' if ok else 'FAIL'
        earned += weight if ok else 0
    row['score'] = earned
    row['severity'] = 'OK' if earned >= 90 else 'WARNING' if earned >= 70 else 'CRITICAL'
    return row
rows = parallel(devices(args.inventory), audit, args.workers)
wb = Workbook()
ws = wb.active
ws.title = 'Compliance'
headers = list(rows[0]) if rows else ['result']
ws.append(headers)
colors = {'PASS': 'C6EFCE', 'FAIL': 'FFC7CE', 'OK': 'C6EFCE', 'WARNING': 'FFEB9C', 'CRITICAL': 'FFC7CE'}
for row in rows:
    ws.append([row.get(h, '') for h in headers])
    for cell in ws[ws.max_row]:
        if str(cell.value) in colors:
            cell.fill = PatternFill('solid', fgColor=colors[str(cell.value)])
for c in ws[1]:
    c.font = Font(bold=True, color='FFFFFF')
    c.fill = PatternFill('solid', fgColor='1F4E78')
ws.freeze_panes = 'A2'
summary = wb.create_sheet('Summary')
counts = Counter((r.get('severity', 'ERROR') for r in rows))
summary.append(['Status', 'Count'])
for k, v in counts.items():
    summary.append([k, v])
path = OUTPUT / f'enterprise_compliance_{stamp()}.xlsx'
wb.save(path)
log.info('Rapport: %s', path)
